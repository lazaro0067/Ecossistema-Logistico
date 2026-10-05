"""Book DPO: padrão operacional por subbloco, com salvamento automático e checklist de aderência.

O checklist é calculado na hora a partir do texto (não é IA): verifica se o padrão
tem responsáveis, frequência, indicador, segurança e passo a passo.
"""
import io
import re

import pandas as pd
import streamlit as st

from config.dpo import STATUS_PADRAO, todos_subblocos
from core import tema, ui
from repositories import dpo_repo

CHECKLIST = [
    ("Responsáveis definidos (RACI)", r"respons|raci|executor|aprovador|dono"),
    ("Frequência / rotina", r"diári|diari|semanal|mensal|turno|frequ|rotina"),
    ("Indicador / meta ligado ao padrão", r"indicador|kpi|meta|%|resultado"),
    ("Segurança (EPI, riscos)", r"seguran|epi|risco|acidente"),
    ("Passo a passo numerado", r"(^|\n)\s*(\d+[\.\)]|-|•)\s+\w"),
    ("Auditoria / verificação", r"audit|verifica|checklist|inspe"),
]


def avaliar(texto: str) -> list[tuple[str, bool]]:
    t = (texto or "").lower()
    return [(nome, bool(re.search(rx, t))) for nome, rx in CHECKLIST] + [("Detalhamento suficiente (300+ caracteres)",
                                                                          len(t.strip()) >= 300)]


def _texto_do_arquivo(arq) -> str:
    nome = arq.name.lower()
    dados = arq.getvalue()
    if nome.endswith(".docx"):
        from docx import Document
        doc = Document(io.BytesIO(dados))
        linhas = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
        for t in doc.tables:
            linhas += [" | ".join(c.text.strip() for c in r.cells if c.text.strip()) for r in t.rows]
        return "\n".join(l for l in linhas if l)
    if nome.endswith((".xlsx", ".xls")):
        return pd.read_excel(io.BytesIO(dados)).to_string(index=False)
    try:
        return dados.decode("utf-8")
    except UnicodeDecodeError:
        return dados.decode("latin-1")


def editor_padrao(operacao_id: int, usuario: dict, modulo: str, subbloco: str, rotulo: str) -> None:
    p = dpo_repo.buscar_padrao(operacao_id, modulo, subbloco) or {}
    k = f"dpo_{modulo}_{subbloco}_{operacao_id}"

    def _salvar():
        dpo_repo.salvar_padrao(operacao_id, modulo, subbloco,
                               st.session_state.get(f"{k}_tit", ""), st.session_state.get(f"{k}_cont", ""),
                               st.session_state.get(f"{k}_resp", ""), st.session_state.get(f"{k}_st", STATUS_PADRAO[0]),
                               usuario["login"])
        ui.avisar("Padrão salvo automaticamente", "success")

    st_atual = p.get("status") if p.get("status") in STATUS_PADRAO else STATUS_PADRAO[0]
    for sufixo, valor in (("tit", p.get("titulo") or f"Padrão DPO — {rotulo}"), ("resp", p.get("responsavel") or ""),
                          ("st", st_atual), ("cont", p.get("conteudo") or "")):
        st.session_state.setdefault(f"{k}_{sufixo}", valor)

    if p.get("dt_atualizacao"):
        st.caption(f"🕒 Atualizado em {p['dt_atualizacao']} por {p.get('atualizado_por') or '—'}")
    c1, c2, c3 = st.columns([3, 2, 1.4])
    c1.text_input("Título do padrão", key=f"{k}_tit", on_change=_salvar)
    c2.text_input("Responsável", key=f"{k}_resp", on_change=_salvar)
    c3.selectbox("Status", STATUS_PADRAO, key=f"{k}_st", on_change=_salvar)

    col_txt, col_chk = st.columns([2.2, 1])
    with col_txt:
        st.text_area("Procedimento operacional (salva ao sair do campo)", height=300,
                     key=f"{k}_cont", on_change=_salvar,
                     placeholder="1. Objetivo\n2. Responsáveis (executor, aprovador)\n3. Passo a passo\n"
                                 "4. Frequência e auditoria\n5. Indicador de resultado\n6. Segurança / EPIs")
        arq = st.file_uploader("Importar de um documento (.docx, .xlsx, .txt)", type=["docx", "xlsx", "xls", "txt", "csv"],
                               key=f"{k}_arq")
        def _importar():
            arquivo = st.session_state.get(f"{k}_arq")
            if not arquivo:
                return
            try:
                st.session_state[f"{k}_cont"] = _texto_do_arquivo(arquivo)
            except Exception as e:
                ui.avisar(f"Não consegui ler o documento: {e}", "error")
            else:
                _salvar()

        if arq:
            st.button("📄 Usar o texto deste documento", key=f"{k}_imp", on_click=_importar)
    with col_chk:
        itens = avaliar(st.session_state.get(f"{k}_cont", ""))
        ok = sum(1 for _, v in itens if v)
        pct = ok / len(itens) * 100
        import pandas as pd

        tema.kpis([{"titulo": "Aderência ao padrão DPO", "valor": ui.pct(pct, 0), "icone": "✅",
                    "status": "bom" if pct >= 85 else "atencao" if pct >= 60 else "critico",
                    "detalhe": f"{ok} de {len(itens)} itens", "ver": "itens que faltam",
                    "dados": pd.DataFrame([{"item": n} for n, v in itens if not v])}], key=f"kp_dpo_{k}")
        st.markdown("\n".join(f"{'✅' if v else '⬜'} {n}" for n, v in itens))


def book(operacao_id: int, usuario: dict, modulo: str, grupo: dict, chave: str,
         extras: dict | None = None) -> None:
    """Desenha um grupo do book (ex.: Fundamentos) — pilar por seletor, subblocos por aba.
    extras = {subbloco: func} para conteúdo adicional (ex.: plantas no 1.1)."""
    pilares = list(grupo)
    pilar = st.selectbox("Pilar", pilares, key=f"pilar_{chave}") if len(pilares) > 1 else pilares[0]
    itens = grupo[pilar]
    if len(itens) == 1:
        sb, rot = itens[0]
        if extras and sb in extras:
            extras[sb]()
        editor_padrao(operacao_id, usuario, modulo, sb, rot)
        return
    for aba, (sb, rot) in zip(st.tabs([r for _, r in itens]), itens):
        with aba:
            if extras and sb in extras:
                extras[sb]()
                st.divider()
            editor_padrao(operacao_id, usuario, modulo, sb, rot)


def progresso_book(operacao_id: int, modulo: str, estrutura: dict, chave: str = "") -> None:
    todos = todos_subblocos(estrutura)
    feitos = dpo_repo.resumo_padroes(operacao_id, modulo)
    status = dict(zip(feitos["subbloco"], feitos["status"])) if not feitos.empty else {}
    impl = sum(1 for sb, _ in todos if status.get(sb) in ("Implementado", "Auditado"))
    escritos = sum(1 for sb, _ in todos if sb in status)
    pct = impl / len(todos) * 100 if todos else 0
    import pandas as pd

    lista = pd.DataFrame([{"subbloco": sb, "padrao": rot, "status": status.get(sb, "Não escrito")} for sb, rot in todos])
    tema.kpis([
        {"titulo": "Padrões implementados", "valor": f"{impl} de {len(todos)}", "icone": "📘",
         "status": "bom" if pct >= 85 else "atencao" if pct >= 50 else "critico", "detalhe": ui.pct(pct, 0),
         "dados": lista[~lista["status"].isin(["Implementado", "Auditado"])], "ver": "o que falta implementar"},
        {"titulo": "Padrões escritos", "valor": escritos, "icone": "✍️", "status": "info",
         "dados": lista[lista["status"] != "Não escrito"]},
        {"titulo": "Auditados", "valor": sum(1 for v in status.values() if v == "Auditado"), "icone": "🔍",
         "status": "info", "dados": lista[lista["status"] == "Auditado"]},
    ], key=f"kp_book_{modulo}_{chave}")
