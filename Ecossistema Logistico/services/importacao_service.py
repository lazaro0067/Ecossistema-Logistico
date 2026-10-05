"""Importação de relatórios (Excel, CSV ou tabela de Word) para o banco.

Cada LAYOUT descreve a tabela de destino e os campos. Para achar a coluna de
cada campo o sistema tenta, nesta ordem:
  1. nome exato (apelidos);
  2. nome que CONTÉM um trecho (ex.: "Disp" acha "Disp." e "Disponível");
  3. POSIÇÃO da coluna no relatório Ambev (ex.: fator HL na coluna Q da 01.11).
Campos com `prioriza_posicao` usam a posição primeiro — é assim que o sistema
original lia os relatórios Ambev (colunas fixas).
"""
import csv
import io
import re
import unicodedata
from dataclasses import dataclass, field

import pandas as pd

from core import tempo
from repositories.base import registrar_log, substituir, upsert
from services import curva_abc_service, produtos_service
from services.erros import RegraNegocioError


@dataclass
class C:
    rotulo: str
    tipo: str = "str"                      # str | int | float | date
    obrigatorio: bool = False
    apelidos: list[str] = field(default_factory=list)
    contem: list[str] = field(default_factory=list)
    posicao: int | None = None             # índice da coluna (A=0, B=1...)
    prioriza_posicao: bool = False


LAYOUTS: dict[str, dict] = {
    "produtos": {
        "rotulo": "Relatório 01.11 — Cadastro de produtos", "frequencia_dias": 180,
        "tabela": "produtos", "chaves": ["cod"], "por_operacao": False, "modo": "upsert",
        "campos": {
            "cod": C("Código", "int", True, ["cod", "codigo", "cod_produto", "codigo_produto", "cod_clean", "sku"],
                     ["codigo", "cod"]),
            "descricao": C("Descrição", apelidos=["descricao", "descricao_produto", "produto"], contem=["descri", "desc"]),
            "fator_hl": C("Fator HL (coluna Q)", "float", apelidos=["fator_hl", "fator_hecto", "hl_cx"], posicao=16,
                          prioriza_posicao=True),
            "cx_pallet": C("Caixas por palete", "float", apelidos=["cx_pallet", "caixas_pallet", "caixas_palete"],
                           contem=["caixas_pallet", "pallet", "palete"]),
            "tipo": C("Tipo", apelidos=["tipo"]),
            "categoria": C("Categoria", apelidos=["categoria"]),
        },
    },
    "linear": {
        "rotulo": "Relatório Linear — média de venda (cx/dia)", "frequencia_dias": 90,
        "tabela": "linear_vendas", "chaves": ["operacao_id", "cod"], "por_operacao": True, "modo": "upsert",
        "campos": {
            "cod": C("Código", "int", True, ["cod", "codigo", "cod_clean", "sku", "item"], ["cod", "item", "produto"],
                     posicao=0),
            "linear_cx_dia": C("Linear (coluna E)", "float", True, ["linear", "linear_vendas", "media_venda"],
                               posicao=4, prioriza_posicao=True),
            "tipo": C("Tipo (Cerveja/NAB)", apelidos=["tipo"]),
            "categoria": C("Categoria", apelidos=["categoria"]),
        },
    },
    "estoque": {
        "rotulo": "Relatório 02.03.04 — Posição de estoque", "frequencia_dias": 1,
        "tabela": "estoque", "chaves": ["operacao_id", "cod"], "por_operacao": True, "modo": "substituir_operacao",
        "campos": {
            "cod": C("Código", "int", True, ["cod", "codigo", "cod_produto", "cod_clean"], ["cod"]),
            "descricao": C("Descrição", apelidos=["descricao"], contem=["desc"]),
            "inicial": C("Inicial", "float", apelidos=["inicial", "estoque_inicial"], contem=["inic"]),
            "entrada": C("Entradas", "float", apelidos=["ent", "entrada", "entradas"], contem=["entr", "ent"]),
            "saida": C("Saídas", "float", apelidos=["saida", "saidas"], contem=["said", "sai"]),
            "disponivel": C("Disponível", "float", True, ["disp", "disponivel", "saldo_disponivel"], ["disp"]),
            "saldo_dia": C("Saldo do dia", "float", apelidos=["saldo_dia"]),
        },
    },
    "pedidos_marcados": {
        "rotulo": "Puxada Marcada (D0, D1, D2)", "frequencia_dias": 1, "data_padrao": "data_puxada",
        "digito_verificador": True,
        "tabela": "pedidos_marcados", "chaves": [], "por_operacao": True, "modo": "substituir_data",
        "campos": {
            "data_puxada": C("Data da puxada", "date", apelidos=["data_puxada", "data"], contem=["data_puxada", "data"]),
            "numero_pedido": C("Nº do pedido", apelidos=["no_pedido", "n_pedido", "numero_pedido", "pedido"],
                               contem=["pedido"]),
            "cod": C("Código (coluna R)", "int", True, ["codigo", "cod"], ["codigo", "cod"], posicao=17,
                     prioriza_posicao=True),
            "descricao": C("Produto", apelidos=["produto", "descricao"], contem=["produto", "desc"]),
            "cx_solicitadas": C("Caixas solicitadas", "float", apelidos=["qtdeskus_item", "solicitado"],
                                contem=["qtdeskus", "solicit"]),
            "cx_marcadas": C("Caixas marcadas (coluna W)", "float", apelidos=["marcado", "cx_marcadas"],
                             contem=["marcad"], posicao=22, prioriza_posicao=True),
            "hl_marcado": C("HL", "float", apelidos=["hl", "hl_marcado"], contem=["hecto", "hl"]),
            "status_item": C("Status", apelidos=["status_item", "status"], contem=["status"]),
        },
    },
    "ressuprimento": {
        "rotulo": "Ressuprimento diário (todas as filiais)", "frequencia_dias": 1,
        "tabela": "ressuprimento_diario", "chaves": ["operacao_id", "data", "cesta"], "por_operacao": True,
        "modo": "multi_operacao",
        "campos": {
            "operacao": C("Operação (coluna A)", apelidos=["operacao", "unidade", "filial"], posicao=0,
                          prioriza_posicao=True),
            "volume_real_hl": C("HL puxado (coluna C)", "float", True, ["hl_puxado", "volume", "sellin"], posicao=2,
                                prioriza_posicao=True),
            "cesta": C("Indicador / cesta (coluna E)", obrigatorio=True, apelidos=["indicador", "cesta"], posicao=4,
                       prioriza_posicao=True),
            "data": C("Data (coluna F)", "date", True, ["data"], posicao=5, prioriza_posicao=True),
        },
    },
    "politica": {
        "rotulo": "Política de estoque (estoque médio / mín / obj / máx)", "frequencia_dias": 7,
        "tabela": "politica_estoque", "chaves": [], "por_operacao": True, "modo": "politica",
        "campos": {
            "data_registro": C("Data (coluna A)", "date", posicao=0, prioriza_posicao=True),
            "sku_original": C("SKU (coluna C)", obrigatorio=True, posicao=2, prioriza_posicao=True),
            "categoria": C("Categoria (coluna D)", posicao=3, prioriza_posicao=True),
            "estoque": C("Estoque (F)", "float", posicao=5, prioriza_posicao=True),
            "demanda": C("Demanda (G)", "float", posicao=6, prioriza_posicao=True),
            "doi_atual": C("DOI atual (H)", "float", posicao=7, prioriza_posicao=True),
            "pe_min_dias": C("Mín. dias (I)", "float", posicao=8, prioriza_posicao=True),
            "pe_obj_dias": C("Obj. dias (J)", "float", posicao=9, prioriza_posicao=True),
            "pe_max_dias": C("Máx. dias (K)", "float", posicao=10, prioriza_posicao=True),
            "pe_min_hl": C("Mín. HL (L)", "float", posicao=11, prioriza_posicao=True),
            "pe_obj_hl": C("Obj. HL (M)", "float", posicao=12, prioriza_posicao=True),
            "pe_max_hl": C("Máx. HL (N)", "float", posicao=13, prioriza_posicao=True),
        },
    },
    "financeiro": {
        "rotulo": "Relatório diário de pagamentos (contas a pagar)", "frequencia_dias": 1,
        "tabela": "contas_pagar", "chaves": [], "por_operacao": True, "modo": "substituir_operacao",
        "campos": {
            "pacote": C("Pacote", apelidos=["pacote"], posicao=0),
            "fornecedor": C("Fornecedor (coluna B)", obrigatorio=True, apelidos=["fornecedor", "nome_fornecedor"],
                            posicao=1, prioriza_posicao=True),
            "departamento": C("Departamento", apelidos=["departamento"], posicao=3),
            "documento": C("Documento / título", apelidos=["documento", "titulo"], posicao=4),
            "historico": C("Histórico", apelidos=["historico"], posicao=7),
            "conta_gerencial": C("Conta gerencial", apelidos=["conta_gerencial"], posicao=8),
            "data_vencimento": C("Vencimento (coluna K)", "date", True, ["vencimento", "data_vencimento"], posicao=10,
                                 prioriza_posicao=True),
            "valor": C("Valor (coluna N)", "float", True, ["valor"], posicao=13, prioriza_posicao=True),
            "realizado": C("Realizado / pago", "float", apelidos=["realizado", "pago", "valor_pago"]),
        },
    },
    "metas_doi": {
        "rotulo": "Metas de DOI por SKU", "frequencia_dias": 90,
        "tabela": "metas_doi", "chaves": ["operacao_id", "cod"], "por_operacao": True, "modo": "upsert",
        "campos": {
            "cod": C("Código", "int", True, ["cod", "codigo", "cod_prod", "cod_clean"], ["cod"]),
            "doi_meta": C("Meta DOI (dias)", "float", True, ["doi", "doi_meta", "meta", "dias"]),
        },
    },
    "curva_abc": {
        "rotulo": "Vendas do mês → Curva ABC", "frequencia_dias": 31,
        "tabela": "curva_abc", "chaves": ["operacao_id", "mes_ano", "cod"], "por_operacao": True,
        "modo": "curva_abc", "pede_mes": True,
        "campos": {
            "cod": C("Código", "int", True, ["cod", "codigo", "cod_clean", "sku"], ["cod"]),
            "descricao": C("Descrição", apelidos=["descricao", "produto"], contem=["desc", "produto"]),
            "total_qtde": C("Quantidade vendida", "float", True, ["qtde", "quantidade", "total", "volume", "total_qtde"],
                            ["qtd", "quant", "volume"]),
        },
    },
}


# --- Leitura do arquivo ----------------------------------------------------
def _norm(txt) -> str:
    txt = unicodedata.normalize("NFKD", str(txt)).encode("ascii", "ignore").decode().lower().strip()
    return re.sub(r"[^a-z0-9]+", "_", txt).strip("_")


def _ler_bruto(nome: str, conteudo) -> pd.DataFrame:
    nome = nome.lower()
    dados = conteudo.read() if hasattr(conteudo, "read") else conteudo
    if nome.endswith((".csv", ".txt")):
        try:
            texto = dados.decode("utf-8-sig")
        except UnicodeDecodeError:
            texto = dados.decode("latin-1")
        linhas = [l for l in texto.splitlines() if l.strip()]
        if not linhas:
            raise ValueError("Arquivo vazio.")
        amostra = "\n".join(linhas[:50])
        sep = max([";", "\t", ",", "|"], key=amostra.count)
        registros = list(csv.reader(linhas, delimiter=sep))
        largura = max(len(r) for r in registros)
        return pd.DataFrame([r + [None] * (largura - len(r)) for r in registros], dtype=object)
    if nome.endswith(".docx"):
        from docx import Document
        doc = Document(io.BytesIO(dados))
        if not doc.tables:
            raise ValueError("O documento Word não tem tabela.")
        tab = max(doc.tables, key=lambda t: len(t.rows))
        return pd.DataFrame([[c.text for c in r.cells] for r in tab.rows], dtype=object)
    return pd.read_excel(io.BytesIO(dados), dtype=str, header=None)


def _pontuar_linha(valores, layouts) -> int:
    nomes = [_norm(v) for v in valores if pd.notna(v) and str(v).strip()]
    pts = 0
    for l in layouts:
        for c in l["campos"].values():
            ap = {_norm(a) for a in c.apelidos}
            if any(n in ap for n in nomes):
                pts += 2
            elif c.contem and any(any(t in n for t in c.contem) for n in nomes):
                pts += 1
    return pts


def ler_arquivo(nome: str, conteudo, layout_key: str | None = None) -> pd.DataFrame:
    """Lê o arquivo e encontra a linha de cabeçalho (relatórios costumam ter títulos acima)."""
    bruto = _ler_bruto(nome, conteudo).dropna(how="all").reset_index(drop=True)
    if bruto.empty:
        raise ValueError("Arquivo sem linhas.")
    layouts = [LAYOUTS[layout_key]] if layout_key else list(LAYOUTS.values())
    melhor, pontos = 0, 0
    for i in range(min(25, len(bruto))):
        p = _pontuar_linha(bruto.iloc[i].tolist(), layouts)
        if p > pontos:
            melhor, pontos = i, p
    cab = [str(c).strip() if pd.notna(c) and str(c).strip() else f"coluna_{j + 1}"
           for j, c in enumerate(bruto.iloc[melhor])]
    vistos: dict[str, int] = {}
    for j, c in enumerate(cab):  # nomes repetidos ganham sufixo
        if c in vistos:
            vistos[c] += 1
            cab[j] = f"{c}_{vistos[c]}"
        else:
            vistos[c] = 0
    df = bruto.iloc[melhor + 1:].copy()
    df.columns = cab
    return df.dropna(how="all").reset_index(drop=True)


def sugerir_mapeamento(layout_key: str, colunas: list[str]) -> dict[str, str | None]:
    campos = LAYOUTS[layout_key]["campos"]
    norm = [_norm(c) for c in colunas]
    usadas: set[str] = set()
    mapa: dict[str, str | None] = {k: None for k in campos}

    def usar(campo, col):
        mapa[campo] = col
        usadas.add(col)

    # 0) posição prioritária (formato fixo dos relatórios Ambev)
    for k, c in campos.items():
        if c.prioriza_posicao and c.posicao is not None and c.posicao < len(colunas):
            usar(k, colunas[c.posicao])
    # 1) nome exato
    for k, c in campos.items():
        if mapa[k]:
            continue
        for a in [k, *c.apelidos]:
            if _norm(a) in norm and colunas[norm.index(_norm(a))] not in usadas:
                usar(k, colunas[norm.index(_norm(a))])
                break
    # 2) nome que contém o trecho
    for k, c in campos.items():
        if mapa[k] or not c.contem:
            continue
        for t in c.contem:
            achou = next((col for col, n in zip(colunas, norm) if t in n and col not in usadas), None)
            if achou:
                usar(k, achou)
                break
    # 3) posição como último recurso
    for k, c in campos.items():
        if not mapa[k] and c.posicao is not None and c.posicao < len(colunas) and colunas[c.posicao] not in usadas:
            usar(k, colunas[c.posicao])
    return mapa


def faltando(layout_key: str, mapa: dict) -> list[str]:
    return [c.rotulo for k, c in LAYOUTS[layout_key]["campos"].items() if c.obrigatorio and not mapa.get(k)]


# --- Conversões --------------------------------------------------------------
def numero_br(v) -> float:
    if v is None or (isinstance(v, float) and pd.isna(v)):
        return 0.0
    if isinstance(v, (int, float)):
        return float(v)
    s = re.sub(r"[R$\s]", "", str(v).strip())
    if not s:
        return 0.0
    if "," in s:  # 1.234,56
        s = s.replace(".", "").replace(",", ".")
    elif re.fullmatch(r"-?[1-9]\d{0,2}(\.\d{3})+", s):  # 1.000 / 12.345.678
        s = s.replace(".", "")
    try:
        return float(s)
    except ValueError:
        return 0.0


_PONTO_DECIMAL = re.compile(r"-?\d+\.\d{1,2}|-?\d+\.\d{4,}|-?\d{4,}\.\d+|-?0\.\d+")


def _numero_ponto(v) -> float:
    """Número com ponto decimal (ex.: 2441.990)."""
    if v is None or (isinstance(v, float) and pd.isna(v)):
        return 0.0
    if isinstance(v, (int, float)):
        return float(v)
    try:
        return float(re.sub(r"[R$\s]", "", str(v)))
    except ValueError:
        return numero_br(v)


def _to_float(s: pd.Series) -> pd.Series:
    """Converte a coluna olhando todos os valores juntos: se nenhum tem vírgula e algum mostra ponto como
    decimal (2441.99 / 2441.990 / 0.5), o ponto é decimal em todos — "441.990" vira 441,99 e não 441 mil."""
    textos = s.dropna().map(lambda v: v if isinstance(v, (int, float)) else str(v).strip())
    textos = textos[textos.map(lambda v: isinstance(v, str) and v != "")]
    limpos = textos.map(lambda v: re.sub(r"[R$\s]", "", v))
    if not limpos.empty and not limpos.str.contains(",").any() and limpos.map(
            lambda v: bool(_PONTO_DECIMAL.fullmatch(v))).any():
        return s.map(_numero_ponto)
    return s.map(numero_br)


def _to_int(s: pd.Series) -> pd.Series:
    return (s.astype(str).str.replace(r"\.0$", "", regex=True).str.replace(r"\D", "", regex=True)
            .replace("", pd.NA).astype("Int64"))


def _to_date(s: pd.Series) -> pd.Series:
    txt = s.astype(str).str.strip()
    iso = txt.str.match(r"^\d{4}-\d{2}-\d{2}")
    out = pd.Series(pd.NaT, index=s.index, dtype="datetime64[ns]")
    out[iso] = pd.to_datetime(txt[iso].str[:10], errors="coerce", format="%Y-%m-%d")
    out[~iso] = pd.to_datetime(txt[~iso], errors="coerce", dayfirst=True)
    return out.dt.strftime("%Y-%m-%d")


def preparar(layout_key: str, df: pd.DataFrame, mapeamento: dict[str, str | None],
             data_padrao: str | None = None, remover_digito: bool = False) -> pd.DataFrame:
    """Aplica o de/para e converte tipos.

    data_padrao: data (AAAA-MM-DD) usada quando o arquivo não traz a coluna de data.
    remover_digito: remove o dígito verificador do código (último dígito), como o original fazia.
    """
    layout = LAYOUTS[layout_key]
    saida = pd.DataFrame(index=df.index)
    for k, c in layout["campos"].items():
        col = mapeamento.get(k)
        if not col:
            if c.obrigatorio:
                raise RegraNegocioError(f"Informe a coluna para o campo obrigatório: {c.rotulo}.")
            continue
        serie = df[col]
        if c.tipo == "int":
            if remover_digito and k == "cod":
                serie = serie.astype(str).str.strip().str.replace(r"\.0$", "", regex=True).str[:-1]
            saida[k] = _to_int(serie)
        elif c.tipo == "float":
            saida[k] = _to_float(serie)
        elif c.tipo == "date":
            saida[k] = _to_date(serie)
        else:
            saida[k] = serie.fillna("").astype(str).str.strip()

    campo_data = layout.get("data_padrao")
    if campo_data and data_padrao:
        saida[campo_data] = saida[campo_data].fillna(data_padrao) if campo_data in saida else data_padrao

    obrig = [k for k, c in layout["campos"].items() if c.obrigatorio]
    if campo_data:
        if campo_data not in saida:
            raise RegraNegocioError("Informe a data (o arquivo não tem coluna de data).")
        obrig.append(campo_data)
    saida = saida.dropna(subset=[o for o in obrig if o in saida])
    if "cod" in saida:
        saida = saida[saida["cod"] > 0]
    for k in ("cesta", "fornecedor", "sku_original"):
        if k in saida:
            saida = saida[saida[k].str.len() > 0]
            saida = saida[~saida[k].str.lower().isin(["nan", "none", "total", "total geral"])]
    if saida.empty:
        raise RegraNegocioError("Nenhuma linha válida após a conversão. Confira o de/para das colunas.")
    return saida


# --- Gravação ------------------------------------------------------------------
def _registros(df: pd.DataFrame) -> list[dict]:
    df = df.astype(object).where(pd.notna(df), None)
    linhas = df.to_dict("records")
    for l in linhas:
        for k, v in l.items():
            if hasattr(v, "item"):
                l[k] = v.item()
    return linhas


def gravar(layout_key: str, dados: pd.DataFrame, operacao_id: int | None, mes_ano: str | None = None,
           usuario: str = "", arquivo: str = "") -> int:
    from repositories import operacoes_repo

    layout = LAYOUTS[layout_key]
    if layout["por_operacao"] and operacao_id and operacoes_repo.e_consolidada(operacao_id) \
            and layout["modo"] != "multi_operacao":
        raise RegraNegocioError("Esta é uma visão consolidada. Escolha uma filial no menu para atualizar a base.")
    df = dados.copy()
    modo, tabela, agora = layout["modo"], layout["tabela"], tempo.agora_str()

    if modo == "curva_abc":
        if not mes_ano:
            raise RegraNegocioError("Informe o mês de referência.")
        if "descricao" not in df:
            df["descricao"] = ""
        df = curva_abc_service.calcular(df)
        df["mes_ano"] = mes_ano

    if modo == "multi_operacao":
        n = _gravar_multi_operacao(df, operacao_id, agora)
        registrar_log(None, layout_key, n, arquivo, usuario)
        return n

    if modo == "politica":
        df["cod"] = df["sku_original"].map(produtos_service.codigo_do_sku)
        df = df[df["cod"] > 0]
        df["tipo"] = df["sku_original"].map(produtos_service.tipo_sku)
        if "data_registro" not in df:
            df["data_registro"] = tempo.hoje().isoformat()
        df["data_registro"] = df["data_registro"].fillna(tempo.hoje().isoformat())

    extras_produto = None
    if layout_key == "linear":
        extras_produto = df[[c for c in ("cod", "tipo", "categoria") if c in df]].copy()
        df = df.drop(columns=[c for c in ("tipo", "categoria") if c in df])

    if layout["por_operacao"]:
        df.insert(0, "operacao_id", operacao_id)
    if tabela in ("estoque", "linear_vendas", "curva_abc", "pedidos_marcados", "politica_estoque", "contas_pagar"):
        df["dt_atualizacao"] = agora
    if layout["chaves"]:
        df = df.drop_duplicates(subset=layout["chaves"], keep="last")
    linhas = _registros(df)

    if modo == "substituir_operacao":
        n = substituir(tabela, "operacao_id = ?", (operacao_id,), linhas, layout["chaves"])
    elif modo in ("substituir_data", "politica"):
        campo = "data_puxada" if modo == "substituir_data" else "data_registro"
        datas = sorted({l[campo] for l in linhas})
        n = substituir(tabela, f"operacao_id = ? AND {campo} IN ({', '.join('?' * len(datas))})",
                       (operacao_id, *datas), linhas, None)
    elif modo == "curva_abc":
        n = substituir(tabela, "operacao_id = ? AND mes_ano = ?", (operacao_id, mes_ano), linhas, None)
    else:
        n = upsert(tabela, linhas, layout["chaves"])

    if extras_produto is not None and len(extras_produto.columns) > 1:
        _atualizar_tipo_categoria(extras_produto)
    registrar_log(operacao_id if layout["por_operacao"] else None, layout_key, n, arquivo, usuario)
    return n


def _gravar_multi_operacao(df: pd.DataFrame, operacao_padrao: int | None, agora: str) -> int:
    """Ressuprimento: um arquivo traz várias filiais; cada linha vai para a sua."""
    from repositories import operacoes_repo

    cache: dict[str, int | None] = {}

    def op_da_linha(txt):
        txt = (txt or "").strip()
        if txt not in cache:
            cache[txt] = operacoes_repo.resolver_por_texto(txt) if txt else None
        if txt:  # filial escrita no arquivo mas não reconhecida: ignora a linha (não joga na filial errada)
            return cache[txt]
        return None if operacoes_repo.e_consolidada(operacao_padrao) else operacao_padrao

    df = df.copy()
    df["operacao_id"] = df["operacao"].map(op_da_linha) if "operacao" in df else operacao_padrao
    sem_op = df["operacao_id"].isna().sum()
    df = df.dropna(subset=["operacao_id"])
    if df.empty:
        raise RegraNegocioError("Não reconheci a filial de nenhuma linha (coluna A). Confira o arquivo.")
    df["operacao_id"] = df["operacao_id"].astype(int)
    df["volume_sellin_hl"] = df["volume_real_hl"]  # o original usa o HL puxado como realizado
    df["dt_atualizacao"] = agora
    df = df.drop(columns=["operacao"], errors="ignore")
    # um valor por revenda + data + indicador (como no sistema original: vale a última linha do arquivo)
    df["cesta"] = df["cesta"].astype(str).str.strip()
    df = df[df["cesta"].ne("") & df["data"].notna()]
    df = df.drop_duplicates(subset=["operacao_id", "data", "cesta"], keep="last")
    # o período que o arquivo traz substitui o que estava gravado (limpa importações erradas)
    from database.connection import execute as _exec
    for op_id, g in df.groupby("operacao_id"):
        _exec("DELETE FROM ressuprimento_diario WHERE operacao_id = ? AND data >= ? AND data <= ?",
              (int(op_id), str(g["data"].min()), str(g["data"].max())))
    n = upsert("ressuprimento_diario", _registros(df), ["operacao_id", "data", "cesta"])
    if sem_op:
        from core import ui  # aviso amigável sem quebrar a importação
        try:
            ui.avisar(f"{sem_op} linha(s) ignoradas: filial não reconhecida na coluna A.", "warning")
        except Exception:
            pass
    return n


def _atualizar_tipo_categoria(df: pd.DataFrame) -> None:
    from database.connection import get_conn

    linhas = [(r.get("tipo") or None, r.get("categoria") or None, r["cod"]) for r in _registros(df)]
    with get_conn() as conn:
        conn.executemany("UPDATE produtos SET tipo = COALESCE(?, tipo), categoria = COALESCE(?, categoria) "
                         "WHERE cod = ?", linhas)
