"""Importação de planilhas (Excel, CSV ou tabela de Word) para o banco.

Cada LAYOUT descreve: tabela de destino, campos esperados e como gravar.
O sistema acha sozinho a linha de cabeçalho, sugere o "de/para" das
colunas pelos apelidos abaixo e, se tudo bater, grava automaticamente.
"""
import csv
import io
import re
import unicodedata

import pandas as pd

from core import tempo
from repositories.base import registrar_log, substituir, upsert
from services import curva_abc_service
from services.erros import RegraNegocioError

# campo: (rótulo, tipo, obrigatório, apelidos de coluna aceitos)
LAYOUTS = {
    "produtos": {
        "rotulo": "Relatório 01.11 — Cadastro de produtos", "frequencia_dias": 180,
        "tabela": "produtos", "chaves": ["cod"], "por_operacao": False, "modo": "upsert",
        "campos": {
            "cod": ("Código", "int", True, ["cod", "codigo", "cod_produto", "codigo_produto", "cod_clean", "sku", "cod_prod"]),
            "descricao": ("Descrição", "str", False, ["descricao", "descricao_produto", "produto", "desc"]),
            "fator_hl": ("Fator HL", "float", False, ["fator_hl", "fator", "fator_hecto", "fator_hectolitro", "hl", "hl_cx"]),
            "cx_pallet": ("Caixas por palete", "float", False, ["cx_pallet", "cx_palete", "caixas_pallet", "caixas_palete", "cx_por_pallet", "pallet", "palete"]),
            "tipo": ("Tipo (Cerveja/NAB...)", "str", False, ["tipo"]),
            "categoria": ("Categoria", "str", False, ["categoria", "embalagem"]),
        },
    },
    "estoque": {
        "rotulo": "Relatório 02.03.04 — Posição de estoque", "frequencia_dias": 1,
        "tabela": "estoque", "chaves": ["operacao_id", "cod"], "por_operacao": True, "modo": "substituir_operacao",
        "campos": {
            "cod": ("Código", "int", True, ["cod", "codigo", "cod_produto", "cod_clean", "produto"]),
            "descricao": ("Descrição", "str", False, ["descricao", "desc"]),
            "inicial": ("Estoque inicial", "float", False, ["inicial", "estoque_inicial"]),
            "entrada": ("Entradas", "float", False, ["entrada", "entradas"]),
            "saida": ("Saídas", "float", False, ["saida", "saidas"]),
            "disponivel": ("Disponível", "float", True, ["disponivel", "saldo_disponivel", "estoque_disponivel", "saldo", "estoque"]),
            "saldo_dia": ("Saldo do dia", "float", False, ["saldo_dia", "saldo dia"]),
        },
    },
    "linear": {
        "rotulo": "Relatório Linear — média de venda (cx/dia)", "frequencia_dias": 90,
        "tabela": "linear_vendas", "chaves": ["operacao_id", "cod"], "por_operacao": True, "modo": "upsert",
        "campos": {
            "cod": ("Código", "int", True, ["cod", "codigo", "cod_clean", "sku"]),
            "linear_cx_dia": ("Linear (cx/dia)", "float", True, ["linear", "linear_vendas", "media", "media_venda", "media_dia", "cx_dia", "venda_media"]),
        },
    },
    "metas_doi": {
        "rotulo": "Metas de DOI por SKU", "frequencia_dias": 90,
        "tabela": "metas_doi", "chaves": ["operacao_id", "cod"], "por_operacao": True, "modo": "upsert",
        "campos": {
            "cod": ("Código", "int", True, ["cod", "codigo", "cod_prod", "cod_clean"]),
            "doi_meta": ("Meta DOI (dias)", "float", True, ["doi", "doi_meta", "meta", "dias"]),
        },
    },
    "curva_abc": {
        "rotulo": "Vendas do mês → Curva ABC", "frequencia_dias": 31,
        "tabela": "curva_abc", "chaves": ["operacao_id", "mes_ano", "cod"], "por_operacao": True,
        "modo": "curva_abc", "pede_mes": True,
        "campos": {
            "cod": ("Código", "int", True, ["cod", "codigo", "cod_clean", "sku"]),
            "descricao": ("Descrição", "str", False, ["descricao", "produto"]),
            "total_qtde": ("Quantidade vendida", "float", True, ["qtde", "quantidade", "total", "volume", "total_qtde"]),
        },
    },
    "pedidos_marcados": {
        "rotulo": "Puxada Marcada (D0, D1, D2)", "frequencia_dias": 1, "data_padrao": "data_puxada",
        "tabela": "pedidos_marcados", "chaves": [], "por_operacao": True, "modo": "substituir_data",
        "campos": {
            "data_puxada": ("Data da puxada", "date", False, ["data", "data_puxada", "dt", "data_entrega", "data_carregamento"]),
            "numero_pedido": ("Nº do pedido", "str", False, ["pedido", "numero_pedido", "n_pedido"]),
            "cod": ("Código", "int", True, ["cod", "codigo", "sku"]),
            "descricao": ("Descrição", "str", False, ["descricao", "produto"]),
            "cx_solicitadas": ("Caixas solicitadas", "float", False, ["solicitado", "cx_solicitadas", "qtd_solicitada"]),
            "cx_marcadas": ("Caixas marcadas", "float", False, ["marcado", "cx_marcadas", "qtd_marcada"]),
            "hl_marcado": ("HL marcado", "float", False, ["hl", "hl_marcado"]),
            "status_item": ("Status", "str", False, ["status", "status_item", "situacao"]),
        },
    },
    "ressuprimento": {
        "rotulo": "Ressuprimento diário por cesta", "frequencia_dias": 1,
        "tabela": "ressuprimento_diario", "chaves": ["operacao_id", "data", "cesta"], "por_operacao": True,
        "modo": "upsert",
        "campos": {
            "data": ("Data", "date", True, ["data", "dia", "data_registro"]),
            "cesta": ("Cesta", "str", True, ["cesta", "categoria"]),
            "volume_sellin_hl": ("Sell-in (HL)", "float", False, ["sellin", "sell_in", "volume_sellin_hl"]),
            "volume_real_hl": ("Real (HL)", "float", True, ["real", "volume_real_hl", "realizado"]),
        },
    },
}


# --- Utilidades de conversão ---------------------------------------------
def _norm(txt: str) -> str:
    txt = unicodedata.normalize("NFKD", str(txt)).encode("ascii", "ignore").decode().lower().strip()
    return re.sub(r"[^a-z0-9]+", "_", txt).strip("_")


def _todos_apelidos(layout_key: str | None) -> set[str]:
    layouts = [LAYOUTS[layout_key]] if layout_key else LAYOUTS.values()
    return {_norm(a) for l in layouts for c, (_, _, _, ap) in l["campos"].items() for a in [c, *ap]}


def _ler_bruto(nome: str, conteudo) -> pd.DataFrame:
    nome = nome.lower()
    dados = conteudo.read() if hasattr(conteudo, "read") else conteudo
    if nome.endswith(".csv") or nome.endswith(".txt"):
        try:
            texto = dados.decode("utf-8-sig")
        except UnicodeDecodeError:
            texto = dados.decode("latin-1")
        linhas = [l for l in texto.splitlines() if l.strip()]
        amostra = "\n".join(linhas[:50])
        sep = max([";", "\t", ",", "|"], key=amostra.count)
        registros = list(csv.reader(linhas, delimiter=sep))
        largura = max(len(r) for r in registros)
        return pd.DataFrame([r + [None] * (largura - len(r)) for r in registros], dtype=object)
    if nome.endswith(".docx"):
        from docx import Document
        doc = Document(io.BytesIO(dados))
        if not doc.tables:
            raise ValueError("O documento Word não tem tabela.")
        tab = max(doc.tables, key=lambda t: len(t.rows))
        return pd.DataFrame([[c.text for c in r.cells] for r in tab.rows], dtype=str)
    return pd.read_excel(io.BytesIO(dados), dtype=str, header=None)


def ler_arquivo(nome: str, conteudo, layout_key: str | None = None) -> pd.DataFrame:
    """Lê o arquivo e encontra a linha de cabeçalho (relatórios costumam ter títulos acima)."""
    bruto = _ler_bruto(nome, conteudo).dropna(how="all").reset_index(drop=True)
    apelidos = _todos_apelidos(layout_key)
    melhor, pontos = 0, -1
    for i in range(min(25, len(bruto))):
        p = sum(_norm(v) in apelidos for v in bruto.iloc[i].fillna(""))
        if p > pontos:
            melhor, pontos = i, p
    cab = [str(c).strip() if pd.notna(c) and str(c).strip() else f"coluna_{j + 1}"
           for j, c in enumerate(bruto.iloc[melhor])]
    vistos: dict[str, int] = {}
    for j, c in enumerate(cab):  # nomes repetidos ganham sufixo
        if c in vistos:
            vistos[c] += 1
            cab[j] = f"{c}_{vistos[c]}"
        else:
            vistos[c] = 0
    df = bruto.iloc[melhor + 1:].copy()
    df.columns = cab
    return df.dropna(how="all").reset_index(drop=True)


def faltando(layout_key: str, mapa: dict) -> list[str]:
    return [rot for c, (rot, _, obrig, _) in LAYOUTS[layout_key]["campos"].items() if obrig and not mapa.get(c)]


def sugerir_mapeamento(layout_key: str, colunas: list[str]) -> dict[str, str | None]:
    normalizadas = {_norm(c): c for c in colunas}
    sugestao = {}
    for campo, (_, _, _, apelidos) in LAYOUTS[layout_key]["campos"].items():
        sugestao[campo] = next((normalizadas[_norm(a)] for a in [campo, *apelidos] if _norm(a) in normalizadas), None)
    return sugestao


def _to_float(s: pd.Series) -> pd.Series:
    def conv(v):
        if v is None or (isinstance(v, float) and pd.isna(v)):
            return 0.0
        v = str(v).strip().replace("R$", "").replace(" ", "")
        if "," in v:  # formato brasileiro 1.234,56
            v = v.replace(".", "").replace(",", ".")
        elif re.fullmatch(r"-?[1-9]\d{0,2}(\.\d{3})+", v):  # 1.000 / 12.345.678 = milhar
            v = v.replace(".", "")
        try:
            return float(v)
        except ValueError:
            return 0.0
    return s.map(conv)


def _to_int(s: pd.Series) -> pd.Series:
    return s.astype(str).str.replace(r"\.0$", "", regex=True).str.replace(r"\D", "", regex=True) \
            .replace("", pd.NA).astype("Int64")


def _to_date(s: pd.Series) -> pd.Series:
    iso = s.astype(str).str.match(r"^\d{4}-\d{2}-\d{2}")
    out = pd.Series(pd.NaT, index=s.index, dtype="datetime64[ns]")
    out[iso] = pd.to_datetime(s[iso].str[:10], errors="coerce", format="%Y-%m-%d")
    out[~iso] = pd.to_datetime(s[~iso], errors="coerce", dayfirst=True)
    return out.dt.strftime("%Y-%m-%d")


def preparar(layout_key: str, df: pd.DataFrame, mapeamento: dict[str, str | None],
             data_padrao: str | None = None, remover_digito: bool = False) -> pd.DataFrame:
    """Aplica o de/para e converte tipos. Levanta erro se faltar campo obrigatório.

    data_padrao: data (AAAA-MM-DD) usada quando o arquivo não traz a coluna de data.
    remover_digito: remove o dígito verificador do código (último dígito).
    """
    layout = LAYOUTS[layout_key]
    saida = pd.DataFrame(index=df.index)
    for campo, (rotulo, tipo, obrig, _) in layout["campos"].items():
        col = mapeamento.get(campo)
        if not col:
            if obrig:
                raise RegraNegocioError(f"Informe a coluna para o campo obrigatório: {rotulo}.")
            continue
        serie = df[col]
        if tipo == "int":
            saida[campo] = _to_int(serie)
        elif tipo == "float":
            saida[campo] = _to_float(serie)
        elif tipo == "date":
            saida[campo] = _to_date(serie)
        else:
            saida[campo] = serie.fillna("").astype(str).str.strip()
    campo_data = layout.get("data_padrao")
    if campo_data and data_padrao:
        if campo_data not in saida:
            saida[campo_data] = data_padrao
        else:
            saida[campo_data] = saida[campo_data].fillna(data_padrao)
    if remover_digito and "cod" in saida:
        saida["cod"] = (saida["cod"] // 10).astype("Int64")
    obrig = [c for c, (_, _, o, _) in layout["campos"].items() if o]
    if campo_data:
        obrig.append(campo_data)
        if campo_data not in saida:
            raise RegraNegocioError("Informe a data (o arquivo não tem coluna de data).")
    saida = saida.dropna(subset=obrig)
    chave_cod = [c for c in ("cod",) if c in saida]
    if chave_cod:
        saida = saida[saida["cod"] > 0]
    if saida.empty:
        raise RegraNegocioError("Nenhuma linha válida após a conversão. Confira o de/para das colunas.")
    return saida


def gravar(layout_key: str, dados: pd.DataFrame, operacao_id: int | None, mes_ano: str | None = None,
           usuario: str = "", arquivo: str = "") -> int:
    layout = LAYOUTS[layout_key]
    df = dados.copy()
    modo = layout["modo"]

    if modo == "curva_abc":
        if not mes_ano:
            raise RegraNegocioError("Informe o mês de referência.")
        if "descricao" not in df:
            df["descricao"] = ""
        df = curva_abc_service.calcular(df)
        df["mes_ano"] = mes_ano

    if layout["por_operacao"]:
        df.insert(0, "operacao_id", operacao_id)
    if layout["tabela"] in ("estoque", "linear_vendas", "curva_abc", "pedidos_marcados", "ressuprimento_diario"):
        df["dt_atualizacao"] = tempo.agora_str()
    if layout["chaves"]:  # duplicados na mesma chave: vale a última linha
        df = df.drop_duplicates(subset=layout["chaves"], keep="last")

    df = df.astype(object).where(pd.notna(df), None)
    linhas = df.to_dict("records")
    for l in linhas:  # tipos numpy/pandas -> nativos
        for k, v in l.items():
            if hasattr(v, "item"):
                l[k] = v.item()

    tabela = layout["tabela"]
    if modo == "substituir_operacao":
        n = substituir(tabela, "operacao_id = ?", (operacao_id,), linhas, layout["chaves"])
    elif modo == "substituir_data":
        datas = sorted({l["data_puxada"] for l in linhas})
        n = substituir(tabela, f"operacao_id = ? AND data_puxada IN ({', '.join('?' * len(datas))})",
                       (operacao_id, *datas), linhas, None)
    elif modo == "curva_abc":
        n = substituir(tabela, "operacao_id = ? AND mes_ano = ?", (operacao_id, mes_ano), linhas, None)
    else:
        n = upsert(tabela, linhas, layout["chaves"])
    registrar_log(operacao_id if layout["por_operacao"] else None, layout_key, n, arquivo, usuario)
    return n
